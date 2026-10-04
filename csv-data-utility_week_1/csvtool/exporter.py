"""Write output files. All writes are atomic (temp file + rename) so a crash never leaves
a half-written report behind. Cell values are neutralised against spreadsheet formula
injection ('=cmd|...' etc.)."""

from __future__ import annotations

import csv
import json
import os
import re
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from typing import IO, Any, Dict, Iterator, List, Sequence

from csvtool.errors import ExportError
from csvtool.parsing import parse_int, parse_money, strip_control
from csvtool.pipeline import PipelineResult, Rejected
from csvtool.reader import Record

_PLAIN_NEGATIVE_NUMBER = re.compile(r"-[0-9]+(?:\.[0-9]+)?")


def safe_cell(value: str) -> str:
    """Prefix risky values with an apostrophe so Excel/Sheets never run them as formulas."""
    text = strip_control(value)
    if text and text[0] in "=+@\t\r" and not _PLAIN_NEGATIVE_NUMBER.fullmatch(text):
        return "'" + text
    if text.startswith("-") and not _PLAIN_NEGATIVE_NUMBER.fullmatch(text):
        return "'" + text
    return text


@contextmanager
def _atomic_text(path: Path) -> Iterator[IO[str]]:
    tmp = path.with_name(path.name + ".tmp")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with tmp.open("w", encoding="utf-8", newline="") as handle:
            yield handle
        os.replace(tmp, path)
    except OSError as exc:
        raise ExportError(f"cannot write {path}: {exc}") from exc
    finally:
        if tmp.exists():
            tmp.unlink()


def write_cleaned_csv(path: Path, columns: Sequence[str], rows: Sequence[Record]) -> None:
    with _atomic_text(path) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(columns)
        for row in rows:
            writer.writerow([safe_cell(row.data.get(c, "")) for c in columns])


def write_rejected_csv(path: Path, columns: Sequence[str], rejected: Sequence[Rejected]) -> None:
    with _atomic_text(path) as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["line", "status", "reason", *columns])
        for item in rejected:
            writer.writerow(
                [
                    item.record.line,
                    item.status,
                    safe_cell(item.reason),
                    *[safe_cell(item.record.data.get(c, "")) for c in columns],
                ]
            )


def write_text(path: Path, text: str) -> None:
    with _atomic_text(path) as handle:
        handle.write(text)


def write_json(path: Path, payload: Dict[str, Any]) -> None:
    with _atomic_text(path) as handle:
        json.dump(payload, handle, indent=2, ensure_ascii=False)
        handle.write("\n")


def _typed(column: str, value: str) -> Any:
    """Give Excel real numbers/dates for the well-known columns."""
    if column == "age" and value:
        age = parse_int(value)
        return age if age is not None else safe_cell(value)
    if column == "salary" and value:
        salary = parse_money(value)
        return salary if salary is not None else safe_cell(value)
    if column == "join_date" and value:
        try:
            return date.fromisoformat(value)
        except ValueError:
            return safe_cell(value)
    return safe_cell(value)


def write_xlsx(
    path: Path,
    result: PipelineResult,
    summary: Dict[str, Any],
) -> None:
    """Excel workbook with Cleaned / Rejected / Summary sheets (needs openpyxl)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ExportError("Excel export needs openpyxl: pip install openpyxl") from exc

    def style(sheet: Any, rows: List[List[Any]]) -> None:
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for index, column in enumerate(sheet.columns, start=1):
            longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 60)

    columns = result.columns
    workbook = Workbook()
    cleaned = workbook.active
    cleaned.title = "Cleaned"
    style(cleaned, [list(columns)] + [[_typed(c, r.data.get(c, "")) for c in columns] for r in result.valid])
    for index, column in enumerate(columns, start=1):
        if column == "join_date":
            for cell in cleaned[get_column_letter(index)][1:]:
                cell.number_format = "yyyy-mm-dd"

    rejected = workbook.create_sheet("Rejected")
    style(
        rejected,
        [["line", "status", "reason", *columns]]
        + [
            [r.record.line, r.status, safe_cell(r.reason)]
            + [safe_cell(r.record.data.get(c, "")) for c in columns]
            for r in result.rejected
        ],
    )

    sheet = workbook.create_sheet("Summary")
    rows: List[List[Any]] = [["metric", "value"]]
    rows += [[k, v] for k, v in summary["counts"].items()]
    for section in ("salary", "age"):
        data = summary["statistics"].get(section)
        if data:
            rows += [[f"{section} {k}", v] for k, v in data.items()]
    for dept, info in summary["statistics"]["by_department"].items():
        rows.append([f"department {dept} (count)", info["count"]])
    style(sheet, rows)

    tmp = path.with_name(path.name + ".tmp")
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(tmp))
        os.replace(tmp, path)
    except OSError as exc:
        raise ExportError(f"cannot write {path}: {exc}") from exc
    finally:
        if tmp.exists():
            tmp.unlink()
